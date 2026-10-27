"""Carga de manuscritos en texto plano, Markdown o Word."""

from __future__ import annotations

from pathlib import Path

try:  # pragma: no cover
    import regex as re
except ImportError:  # pragma: no cover
    import re  # type: ignore[no-redef]

from onirifix.modelo import Capitulo, Documento

SUFIJOS_TEXTO = {".txt", ".md", ".markdown", ".rst", ".text"}
SUFIJOS_DOCX = {".docx"}

_PATRON_CAPITULO = re.compile(
    r"^\s*(?:"
    r"cap[ií]tulo|cap\.|parte|libro|acto|escena|prólogo|prologo|epílogo|epilogo|"
    r"introducción|introduccion|prefacio|apéndice|apendice"
    r")\b",
    flags=re.IGNORECASE | re.UNICODE,
)
_PATRON_MARKDOWN = re.compile(r"^\s*#{1,3}\s+\S")


class ErrorLectura(ValueError):
    pass


def sufijo_soportado(ruta: Path) -> bool:
    return ruta.suffix.lower() in SUFIJOS_TEXTO | SUFIJOS_DOCX


def leer_texto(ruta: Path) -> tuple[str, str]:
    ruta = ruta.expanduser()
    if not ruta.exists():
        raise ErrorLectura(f"No existe el archivo: {ruta}")
    if ruta.suffix.lower() in SUFIJOS_DOCX:
        return _leer_docx(ruta), "docx"
    if ruta.suffix.lower() not in SUFIJOS_TEXTO:
        raise ErrorLectura(
            f"Formato no soportado: {ruta.suffix or '(sin extensión)'}. "
            "Usa .txt, .md o .docx."
        )

    datos = ruta.read_bytes()
    if datos.startswith(b"\xef\xbb\xbf"):
        texto = datos.decode("utf-8-sig")
        encoding = "utf-8"
    else:
        for encoding in ("utf-8", "cp1252", "latin-1"):
            try:
                texto = datos.decode(encoding)
                break
            except UnicodeDecodeError:
                continue
        else:  # pragma: no cover - latin-1 nunca falla
            raise ErrorLectura(f"No se pudo decodificar {ruta}.")

    texto = texto.replace("\r\n", "\n").replace("\r", "\n")
    return texto, encoding


def _leer_docx(ruta: Path) -> str:
    try:
        from docx import Document  # type: ignore[import-not-found]
    except ImportError as exc:  # pragma: no cover
        raise ErrorLectura(
            "Para leer .docx instala el extra: pip install -e '.[docx]'"
        ) from exc

    documento = Document(str(ruta))
    bloques: list[str] = []
    for parrafo in documento.paragraphs:
        bloques.append(parrafo.text)
    for tabla in documento.tables:
        for fila in tabla.rows:
            celdas = [celda.text.strip() for celda in fila.cells]
            if any(celdas):
                bloques.append(" | ".join(celdas))
    texto = "\n".join(bloques)
    return texto.replace("\r\n", "\n").replace("\r", "\n")


def detectar_capitulos(texto: str) -> tuple[Capitulo, ...]:
    """Detecta capítulos/escenas por encabezados o líneas canónicas."""

    inicios: list[tuple[int, int, str]] = []
    posicion = 0
    lineas = texto.splitlines(keepends=True)
    # Si el manuscrito usa Markdown, los encabezados manda: una línea de prosa
    # que empiece por «Prefacio» no debe convertirse en un capítulo falso.
    hay_markdown = any(_PATRON_MARKDOWN.match(linea.rstrip("\n")) for linea in lineas)
    for numero_linea, contenido_linea in enumerate(lineas, start=1):
        contenido = contenido_linea.rstrip("\n")
        es_encabezado = (
            bool(_PATRON_MARKDOWN.match(contenido))
            if hay_markdown
            else bool(_PATRON_CAPITULO.match(contenido))
        )
        if es_encabezado:
            inicios.append((posicion, numero_linea, contenido.strip()))
        posicion += len(contenido_linea)

    if not inicios:
        return (Capitulo(titulo="Texto completo", inicio=0, fin=len(texto), linea=1),)

    capitulos: list[Capitulo] = []
    if inicios[0][0] > 0 and texto[: inicios[0][0]].strip():
        capitulos.append(
            Capitulo(titulo="Preámbulo", inicio=0, fin=inicios[0][0], linea=1)
        )
    for indice, (inicio, numero_linea, titulo) in enumerate(inicios):
        fin = inicios[indice + 1][0] if indice + 1 < len(inicios) else len(texto)
        capitulos.append(
            Capitulo(
                titulo=titulo.lstrip("#").strip() or f"Capítulo {indice + 1}",
                inicio=inicio,
                fin=fin,
                linea=numero_linea,
            )
        )
    return tuple(capitulos)


def cargar_documento(ruta: Path) -> Documento:
    texto, encoding = leer_texto(ruta)
    return Documento(
        ruta=ruta,
        texto=texto,
        encoding=encoding,
        capitulos=detectar_capitulos(texto),
    )


def escribir_texto(ruta: Path, texto: str, encoding: str = "utf-8") -> Path:
    ruta = ruta.expanduser()
    if encoding == "docx":
        raise ErrorLectura("La escritura .docx todavía no está implementada; exporta a .md/.txt.")
    ruta.parent.mkdir(parents=True, exist_ok=True)
    with ruta.open("w", encoding=encoding, newline="\n") as archivo:
        archivo.write(texto)
    return ruta
